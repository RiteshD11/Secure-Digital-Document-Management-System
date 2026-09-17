from __future__ import annotations

import re
import csv
import io
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path

import fitz
from docx import Document
from openpyxl import load_workbook
from PIL import Image
from pptx import Presentation

from app.config import get_settings
from app.services.ocr_service import OCRService
from app.utils.logging import get_logger

logger = get_logger(__name__)


class DocumentTextExtractor:
    def __init__(self) -> None:
        self.settings = get_settings()
        self.allowed_extensions = {ext.strip().lower() for ext in self.settings.allowed_file_extensions.split(",") if ext.strip()}
        self.ocr_service = OCRService()
        self.source_type = ""

    _IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp", ".tif", ".tiff"}
    _DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".doc", ".txt", ".csv", ".xlsx", ".pptx"}
    _MIME_TYPES = {
        ".pdf": {"application/pdf"},
        ".docx": {"application/vnd.openxmlformats-officedocument.wordprocessingml.document"},
        ".doc": {"application/msword", "application/vnd.ms-office"},
        ".txt": {"text/plain"},
        ".csv": {"text/csv", "application/csv", "text/plain"},
        ".xlsx": {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
        ".pptx": {"application/vnd.openxmlformats-officedocument.presentationml.presentation"},
        ".jpg": {"image/jpeg"},
        ".jpeg": {"image/jpeg"},
        ".png": {"image/png"},
        ".webp": {"image/webp"},
        ".tif": {"image/tiff"},
        ".tiff": {"image/tiff"},
    }

    def is_supported_file(self, filename: str) -> bool:
        suffix = Path(filename).suffix.lower()
        return suffix in self.allowed_extensions

    def validate_file(self, file_bytes: bytes, filename: str, content_type: str | None = None) -> str:
        if not self.is_supported_file(filename):
            raise ValueError("Unsupported file type.")

        suffix = Path(filename).suffix.lower()
        normalized_content_type = (content_type or "").split(";", 1)[0].strip().lower()
        if normalized_content_type and normalized_content_type not in {"application/octet-stream", "binary/octet-stream"}:
            if normalized_content_type not in self._MIME_TYPES[suffix]:
                raise ValueError("Uploaded content type does not match the file extension.")

        actual_type = self._detect_file_type(file_bytes)
        expected_type = self._source_type_for_extension(suffix)
        if not self._matches_expected_type(actual_type, suffix):
            raise ValueError("Uploaded file content does not match the file extension.")
        return expected_type

    def clean_text(self, text: str) -> str:
        text = re.sub(r"\s+", " ", text)
        text = re.sub(r"\n+", " ", text)
        text = re.sub(r"[\t\r\f]+", " ", text)
        text = text.strip()
        return text

    async def extract_text(self, file_bytes: bytes, filename: str) -> str:
        file_ext = Path(filename).suffix.lower()
        self.source_type = self.validate_file(file_bytes, filename)

        if file_ext == ".pdf":
            return self._extract_pdf_text(file_bytes)

        if file_ext == ".docx":
            return self._extract_docx_text(file_bytes)

        if file_ext == ".doc":
            return self._extract_doc_text(file_bytes)

        if file_ext == ".txt":
            return self._decode_text(file_bytes)

        if file_ext == ".csv":
            return self._extract_csv_text(file_bytes)

        if file_ext == ".xlsx":
            return self._extract_xlsx_text(file_bytes)

        if file_ext == ".pptx":
            return self._extract_pptx_text(file_bytes)

        if file_ext in self._IMAGE_EXTENSIONS:
            return self._extract_image_text(file_bytes)

        raise ValueError("Unsupported file type.")

    @staticmethod
    def _source_type_for_extension(file_ext: str) -> str:
        return file_ext.lstrip(".")

    @staticmethod
    def _matches_expected_type(actual_type: str, file_ext: str) -> bool:
        if file_ext in {".txt", ".csv"}:
            return actual_type == "text"
        if file_ext == ".jpg":
            return actual_type in {"jpg", "jpeg"}
        if file_ext == ".jpeg":
            return actual_type in {"jpg", "jpeg"}
        if file_ext in {".tif", ".tiff"}:
            return actual_type == "tiff"
        return actual_type == file_ext.lstrip(".")

    def _detect_file_type(self, file_bytes: bytes) -> str:
        if not file_bytes:
            raise ValueError("Uploaded file is empty.")
        if file_bytes.startswith(b"%PDF-"):
            return "pdf"
        if file_bytes.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
            return "doc"
        if file_bytes.startswith(b"PK\x03\x04"):
            try:
                with zipfile.ZipFile(io.BytesIO(file_bytes)) as archive:
                    members = set(archive.namelist())
            except (OSError, zipfile.BadZipFile) as exc:
                raise ValueError("Uploaded Office document is corrupted.") from exc
            if "word/document.xml" in members:
                return "docx"
            if "xl/workbook.xml" in members:
                return "xlsx"
            if "ppt/presentation.xml" in members:
                return "pptx"
            raise ValueError("Unsupported Office document type.")

        try:
            with Image.open(io.BytesIO(file_bytes)) as image:
                image.verify()
                image_type = image.format.lower()
        except Exception as exc:
            image_type = ""
        image_types = {"jpg": "jpg", "jpeg": "jpeg", "png": "png", "webp": "webp", "tiff": "tiff"}
        if image_type in image_types:
            return image_types[image_type]

        try:
            self._decode_text(file_bytes)
        except ValueError:
            raise ValueError("Unable to detect the uploaded file type.") from None
        return "text"

    def _extract_pdf_text(self, file_bytes: bytes) -> str:
        try:
            with fitz.open(stream=file_bytes, filetype="pdf") as pdf:
                text_parts = []
                for page_number, page in enumerate(pdf, start=1):
                    page_text = page.get_text("text")
                    if page_text:
                        text_parts.append(f"Page {page_number}\n{page_text}")
                extracted = "\n".join(text_parts)
                if len(extracted.strip()) >= 20:
                    return extracted

                ocr_parts = []
                for page_number, page in enumerate(pdf, start=1):
                    image_bytes = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False).tobytes("png")
                    page_text = self.ocr_service.extract_text_from_image(image_bytes)
                    if page_text.strip():
                        ocr_parts.append(f"Page {page_number}\n{page_text}")
                if ocr_parts:
                    return "\n".join(ocr_parts)
                raise ValueError("PDF does not contain extractable text.")
        except Exception as exc:
            if isinstance(exc, ValueError) and str(exc) == "PDF does not contain extractable text.":
                raise
            logger.warning("PDF extraction failed for stream-based extraction: %s", exc)
            raise ValueError("Unable to extract text from the PDF.") from exc

    def _extract_docx_text(self, file_bytes: bytes) -> str:
        try:
            document = Document(io.BytesIO(file_bytes))
            parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                parts.extend(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
            return "\n".join(parts)
        except Exception as exc:
            raise ValueError("Unable to extract text from the DOCX document.") from exc

    def _extract_doc_text(self, file_bytes: bytes) -> str:
        converter = shutil.which("soffice") or shutil.which("libreoffice")
        if not converter:
            raise ValueError("Legacy DOC processing requires LibreOffice, which is unavailable.")
        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "uploaded.doc"
            source_path.write_bytes(file_bytes)
            try:
                result = subprocess.run(
                    [converter, "--headless", "--convert-to", "docx", "--outdir", temp_dir, str(source_path)],
                    capture_output=True,
                    timeout=30,
                    check=False,
                )
            except subprocess.TimeoutExpired as exc:
                raise ValueError("Legacy DOC conversion timed out.") from exc
            converted_path = Path(temp_dir) / "uploaded.docx"
            if result.returncode != 0 or not converted_path.exists():
                raise ValueError("Unable to convert the legacy DOC document for processing.")
            return self._extract_docx_text(converted_path.read_bytes())

    def _decode_text(self, file_bytes: bytes) -> str:
        for encoding in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
            try:
                text = file_bytes.decode(encoding)
                if "\x00" not in text:
                    return text
            except UnicodeDecodeError:
                continue
        raise ValueError("Unable to decode the text document.")

    def _extract_csv_text(self, file_bytes: bytes) -> str:
        try:
            text = self._decode_text(file_bytes)
            rows = csv.reader(io.StringIO(text))
            return "\n".join(" | ".join(cell.strip() for cell in row) for row in rows)
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("Unable to extract text from the CSV document.") from exc

    def _extract_xlsx_text(self, file_bytes: bytes) -> str:
        try:
            workbook = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
            parts = []
            for worksheet in workbook.worksheets:
                parts.append(f"Sheet: {worksheet.title}")
                for row in worksheet.iter_rows(values_only=True):
                    values = [str(value).strip() for value in row if value is not None and str(value).strip()]
                    if values:
                        parts.append(" | ".join(values))
            workbook.close()
            return "\n".join(parts)
        except Exception as exc:
            raise ValueError("Unable to extract text from the XLSX document.") from exc

    def _extract_pptx_text(self, file_bytes: bytes) -> str:
        try:
            presentation = Presentation(io.BytesIO(file_bytes))
            parts = []
            for slide_number, slide in enumerate(presentation.slides, start=1):
                slide_parts = []
                for shape in slide.shapes:
                    if getattr(shape, "has_text_frame", False) and shape.text.strip():
                        slide_parts.append(shape.text.strip())
                    if getattr(shape, "has_table", False):
                        slide_parts.extend(
                            " | ".join(cell.text.strip() for cell in row.cells)
                            for row in shape.table.rows
                        )
                if slide_parts:
                    parts.append(f"Slide {slide_number}\n" + "\n".join(slide_parts))
            return "\n".join(parts)
        except Exception as exc:
            raise ValueError("Unable to extract text from the PPTX presentation.") from exc

    def _extract_image_text(self, file_bytes: bytes) -> str:
        extracted = self.ocr_service.extract_text_from_image(file_bytes)
        if not extracted.strip():
            raise ValueError("OCR could not extract readable text from the image.")
        return extracted
