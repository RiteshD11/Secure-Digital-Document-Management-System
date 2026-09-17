from __future__ import annotations

import asyncio
import io
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock

import fitz
import pytest
from docx import Document
from fastapi import HTTPException, UploadFile
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from app.api.routes.indexing import index_document
from app.services.text_extractor import DocumentTextExtractor


def make_pdf(*page_texts: str) -> bytes:
    document = fitz.open()
    for page_text in page_texts:
        page = document.new_page()
        page.insert_text((72, 72), page_text)
    content = document.tobytes()
    document.close()
    return content


def make_docx() -> bytes:
    document = Document()
    document.add_paragraph("DOCX searchable paragraph")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Case"
    table.cell(0, 1).text = "DOCX-001"
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def make_xlsx() -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Evidence"
    worksheet.append(["Item", "Value"])
    worksheet.append(["XLSX searchable cell", "XLSX-001"])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def make_pptx() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "PPTX searchable title"
    textbox = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(5), Inches(1))
    textbox.text = "PPTX searchable body"
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


def make_image(format_name: str) -> bytes:
    output = io.BytesIO()
    Image.new("RGB", (40, 40), color="white").save(output, format=format_name)
    return output.getvalue()


@pytest.mark.parametrize(
    ("filename", "content_factory", "expected"),
    [
        ("sample.pdf", lambda: make_pdf("PDF searchable text"), "PDF searchable text"),
        ("sample.docx", make_docx, "DOCX searchable paragraph"),
        ("sample.txt", lambda: "TXT searchable text".encode(), "TXT searchable text"),
        ("sample.csv", lambda: b"Header,Value\nCSV searchable cell,CSV-001\n", "CSV searchable cell"),
        ("sample.xlsx", make_xlsx, "XLSX searchable cell"),
        ("sample.pptx", make_pptx, "PPTX searchable title"),
    ],
)
def test_document_formats_extract_searchable_text(filename, content_factory, expected):
    extractor = DocumentTextExtractor()
    text = asyncio.run(extractor.extract_text(content_factory(), filename))
    assert expected in text
    assert extractor.source_type == filename.rsplit(".", 1)[1]


@pytest.mark.parametrize(
    ("filename", "format_name", "expected_source"),
    [
        ("sample.jpg", "JPEG", "jpg"),
        ("sample.png", "PNG", "png"),
        ("sample.webp", "WEBP", "webp"),
        ("sample.tiff", "TIFF", "tiff"),
    ],
)
def test_image_formats_use_ocr(monkeypatch, filename, format_name, expected_source):
    extractor = DocumentTextExtractor()
    monkeypatch.setattr(extractor.ocr_service, "extract_text_from_image", lambda _: "OCR searchable text")

    text = asyncio.run(extractor.extract_text(make_image(format_name), filename))

    assert text == "OCR searchable text"
    assert extractor.source_type == expected_source


def test_multipage_pdf_preserves_page_information():
    extractor = DocumentTextExtractor()
    text = asyncio.run(extractor.extract_text(make_pdf("first page", "second page"), "multi.pdf"))

    assert "Page 1" in text
    assert "Page 2" in text
    assert "second page" in text


def test_scanned_pdf_uses_ocr(monkeypatch):
    extractor = DocumentTextExtractor()
    ocr_calls = []

    def fake_ocr(image_bytes):
        ocr_calls.append(image_bytes)
        return "Scanned PDF text"

    monkeypatch.setattr(extractor.ocr_service, "extract_text_from_image", fake_ocr)
    document = fitz.open()
    document.new_page()
    scanned_pdf = document.tobytes()
    document.close()

    text = asyncio.run(extractor.extract_text(scanned_pdf, "scanned.pdf"))

    assert "Scanned PDF text" in text
    assert ocr_calls


def test_unsupported_extension_is_rejected():
    extractor = DocumentTextExtractor()

    with pytest.raises(ValueError, match="Unsupported file type"):
        extractor.validate_file(b"#!/bin/sh\necho unsafe", "script.sh", "text/x-shellscript")


def test_mismatched_content_type_is_rejected():
    extractor = DocumentTextExtractor()

    with pytest.raises(ValueError, match="does not match"):
        extractor.validate_file(make_image("PNG"), "image.png", "application/pdf")


def test_corrupted_file_is_rejected():
    extractor = DocumentTextExtractor()

    with pytest.raises(ValueError, match="does not match|detect"):
        extractor.validate_file(b"not a PDF", "broken.pdf", "application/pdf")


def test_legacy_doc_reports_missing_converter(monkeypatch):
    extractor = DocumentTextExtractor()
    monkeypatch.setattr("app.services.text_extractor.shutil.which", lambda _: None)
    legacy_doc_header = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"legacy doc"

    with pytest.raises(ValueError, match="LibreOffice"):
        asyncio.run(extractor.extract_text(legacy_doc_header, "legacy.doc"))


def test_empty_document_returns_http_400():
    request = MagicMock()
    request.app.state.embedding_service = MagicMock()
    request.app.state.vector_service = MagicMock()
    file = UploadFile(filename="empty.txt", file=io.BytesIO(b""))

    with pytest.raises(HTTPException) as error:
        asyncio.run(index_document(request, file=file, document_id="DOC1", case_id="CASE1"))

    assert error.value.status_code == 400
    assert "empty" in error.value.detail.lower()


def test_oversized_file_returns_http_400():
    request = MagicMock()
    request.app.state.embedding_service = MagicMock()
    request.app.state.vector_service = MagicMock()
    file = UploadFile(filename="large.txt", file=io.BytesIO(b"x" * (20 * 1024 * 1024 + 1)))

    with pytest.raises(HTTPException) as error:
        asyncio.run(index_document(request, file=file, document_id="DOC1", case_id="CASE1"))

    assert error.value.status_code == 400
    assert "maximum" in error.value.detail.lower()


@pytest.mark.asyncio
async def test_upload_stores_source_type_without_document_content():
    request = MagicMock()
    request.app.state.embedding_service = SimpleNamespace(embed_texts=AsyncMock(return_value=[[0.1, 0.2]]))
    request.app.state.vector_service = SimpleNamespace(
        ensure_collection=AsyncMock(),
        delete_document_points=AsyncMock(),
        upsert_points=AsyncMock(),
    )
    file = UploadFile(filename="case.txt", file=io.BytesIO(b"secret searchable text"))

    response = await index_document(request, file=file, document_id="DOC1", case_id="CASE1")

    assert response.success is True
    point = request.app.state.vector_service.upsert_points.await_args.args[0][0]
    assert point["payload"]["source_type"] == "txt"
    assert "secret searchable text" not in point["payload"]
    assert "embedding" not in point["payload"]
